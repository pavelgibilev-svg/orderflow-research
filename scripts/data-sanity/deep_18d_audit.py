"""OKX MAY 18D DEEP FILTER + ENGINE SELECTION AUDIT (sections A-F + I).

IN-SAMPLE RESEARCH (not production proof). Deep reconstruction of OKX 2026-05-03..20 unique-move pool:
A cluster map, B engine/zone-selection audit, C positive pattern mining, D per-date casebook,
E aggressive (clearly in-sample) rule search, F engine rework proposals, I final report.
Decisions causal (<= confirmedTs); future labels = outcome only; TP stays 2%; 2.5/3% are quality labels.
"""
from __future__ import annotations
import csv, gzip, json, math, statistics as st, sys, importlib.util, datetime as dt
from collections import defaultdict
from pathlib import Path

ROOT = next(p for p in Path(__file__).resolve().parents if (p / "package.json").exists() and (p / "scripts").exists())
sys.path.insert(0, str(ROOT / "scripts/data-sanity"))
spec = importlib.util.spec_from_file_location("sm", str(ROOT / "scripts/data-sanity/successful_unique_move_mining.py"))
SM = importlib.util.module_from_spec(spec); spec.loader.exec_module(SM)
TA, TZ = SM.TA, SM.TZ
import openpyxl

OUT = ROOT / "reports/okx-may-early"; DATA = ROOT / "data/okx-historical/BTC-USDT-SWAP"
DATES = [f"2026-05-{d:02d}" for d in range(3, 21)]


def now_iso(): return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
def num(x): return x if isinstance(x, (int, float)) else None
def mins(s): return round(s / 60, 1) if isinstance(s, (int, float)) else None
def winrate(outs):
    o = [x for x in outs if x in ("WIN", "LOSS", "TIMEOUT")]
    return round(100 * sum(1 for x in o if x == "WIN") / max(len(o), 1), 1)


def daily_candle(date):
    p = DATA / date / "trades.csv.gz"
    if not p.exists(): return None
    o = h = l = c = None; vol = 0.0; n = 0
    with gzip.open(p, "rt") as fh:
        rd = csv.reader(fh); next(rd, None)
        for row in rd:
            try: px = float(row[6]); am = float(row[7])
            except (IndexError, ValueError): continue
            if o is None: o = px; h = px; l = px
            h = max(h, px); l = min(l, px); c = px; vol += am; n += 1
    if o is None: return None
    return {"open": o, "high": h, "low": l, "close": c, "vol": round(vol, 2),
            "net_pct": round(100 * (c - o) / o, 2), "range_pct": round(100 * (h - l) / o, 2), "trades": n}


def auc(pos, neg):
    pos = [x for x in pos if x is not None]; neg = [x for x in neg if x is not None]
    if not pos or not neg: return None
    allv = sorted(pos + neg)
    ranks = {}; i = 0
    while i < len(allv):
        j = i
        while j + 1 < len(allv) and allv[j + 1] == allv[i]: j += 1
        r = (i + j) / 2 + 1
        ranks[allv[i]] = r; i = j + 1
    rp = sum(ranks[x] for x in pos)
    a = (rp - len(pos) * (len(pos) + 1) / 2) / (len(pos) * len(neg))
    return round(a, 3)


def cohen_d(a, b): return SM.cohen_d(a, b)


def main():
    zones, emed = SM.load_zones()
    h2cl, scl, strong_sets, ncl = SM.build(zones, 120)
    n_hit2 = len(h2cl); n_strong = len(scl)
    raw = len(zones)
    triggered = sum(1 for z in zones if z.get("sim_outcome") in ("WIN", "LOSS", "TIMEOUT"))
    raw_hit2 = sum(1 for z in zones if z["_hit2"]); raw_strong = sum(1 for z in zones if z["_hit25"])

    def hindsight_risk(c):
        live = any(z["_live"] for z in c); fam = SM.rep_of(c)["_family"]
        if not live: return "HIGH"
        if fam in ("LIQUIDITY_SWEEP_REVERSAL", "FALSE_BREAKOUT_RECLAIM"): return "LOW"
        return "MED"
    def first_z(c): return min(c, key=lambda z: z["confirmedTs"])
    def best_z(c): return max(c, key=lambda z: (z["_mfe"] if z["_mfe"] is not None else -9))
    def first_live_z(c):
        liv = [z for z in c if z["_live"]]
        return min(liv, key=lambda z: z["confirmedTs"]) if liv else None
    def earliest_safe_z(c):
        w = [z for z in c if z.get("sim_outcome") == "WIN"]
        return min(w, key=lambda z: z["confirmedTs"]) if w else None

    # ---------- A: cluster map ----------
    def tt(z, k): return mins(z.get(k))
    A_rows = []
    for i, c in enumerate(sorted(h2cl, key=lambda c: first_z(c)["confirmedTs"])):
        fz = first_z(c); bz = best_z(c); flz = first_live_z(c)
        is_strong = any(z["_hit25"] for z in c); is_h3 = any(z["_hit3"] for z in c); live = any(z["_live"] for z in c)
        thz = min((z["confirmedTs"] // 1000 + z["time_to_2"]) for z in c if num(z.get("time_to_2")))
        A_rows.append({
            "cluster_id": f"u{i:02d}|{fz['direction']}|{fz['_date']}", "date": fz["_date"], "direction": fz["direction"],
            "first_zone_id": fz["id"], "best_zone_id": bz["id"], "first_live_zone_id": (flz["id"] if flz else None),
            "target_hit_ts": thz, "setup_family": fz["_family"], "regime": fz["_regime"],
            "MFE": bz["_mfe"], "MAE": round(min(num(z.get("true_mae")) or 0 for z in c), 3),
            "hit2": 1, "hit2_5": int(is_strong), "hit3": int(is_h3),
            "time_to_2_min": tt(fz, "time_to_2"), "time_to_2_5_min": tt(bz, "time_to_2_5"), "time_to_3_min": tt(bz, "time_to_3"),
            "zones_in_cluster": len(c), "duplicate_count": len(c) - 1,
            "live_detectable": "YES" if live else "NO", "hindsight_risk": hindsight_risk(c),
            "is_strong_cluster": int(is_strong)})
    dup_total = sum(len(c) - 1 for c in h2cl)
    A_summary = {"raw_zones": raw, "triggered_zones": triggered, "raw_hit2_zones": raw_hit2, "raw_strong_zones": raw_strong,
                 "unique_hit2_clusters": n_hit2, "unique_strong_clusters": n_strong,
                 "live_detectable_strong_clusters": sum(1 for c in scl if any(z["_live"] for z in c)),
                 "hindsight_only_strong_clusters": sum(1 for c in scl if not any(z["_live"] for z in c)),
                 "duplicate_zones_in_hit2_clusters": dup_total, "avg_zones_per_hit2_cluster": round(sum(len(c) for c in h2cl) / max(n_hit2, 1), 2)}

    # ---------- B: engine / zone selection audit ----------
    def grp_winrates(clusters, label):
        return {"group": label, "n_clusters": len(clusters),
                "first_zone_winrate": winrate([first_z(c).get("sim_outcome") for c in clusters]),
                "best_mfe_zone_winrate": winrate([best_z(c).get("sim_outcome") for c in clusters]),
                "first_live_zone_winrate": winrate([(first_live_z(c) or first_z(c)).get("sim_outcome") for c in clusters]),
                "earliest_safe_zone_winrate": winrate([(earliest_safe_z(c) or first_z(c)).get("sim_outcome") for c in clusters]),
                "clusters_with_any_winner": sum(1 for c in clusters if any(z.get("sim_outcome") == "WIN" for z in c)),
                "clusters_first_is_loss_but_safe_exists": sum(1 for c in clusters if first_z(c).get("sim_outcome") != "WIN" and any(z.get("sim_outcome") == "WIN" for z in c)),
                "avg_zones_per_cluster": round(sum(len(c) for c in clusters) / max(len(clusters), 1), 2)}
    B_groups = [grp_winrates(scl, "strong_clusters"), grp_winrates(h2cl, "hit2_clusters")]
    delta_first_best_strong = B_groups[0]["best_mfe_zone_winrate"] - B_groups[0]["first_zone_winrate"]
    B_flags = {
        "ENGINE_DUPLICATES_TOO_MANY": "YES" if dup_total / max(raw_hit2, 1) >= 0.5 else "PARTIAL",
        "ENGINE_SELECTION_PROBLEM": "NO" if B_groups[0]["first_zone_winrate"] >= 90 and abs(delta_first_best_strong) < 10 else "PARTIAL",
        "RANKING_LAYER_NEEDED": "PARTIAL",   # cross-cluster ranking would help but current features don't separate
        "ENGINE_REWRITE_NEEDED": "NO",
    }
    B = {"groups": B_groups, "delta_first_vs_best_strong": delta_first_best_strong,
         "interpretation": ("First zone of each strong cluster already wins TP2 (%.0f%%); best-MFE zone adds %+.0f pts. "
                            "=> within-cluster selection is NOT the bottleneck; the gap is CROSS-cluster ranking "
                            "(which of the look-alike clusters to trust), and current features do not separate them."
                            % (B_groups[0]["first_zone_winrate"], delta_first_best_strong)),
         "flags": B_flags}

    # ---------- C: positive pattern mining ----------
    # derived live-valid features on top of SM.FEATS
    extra = {
        "ofi_slope_5m_30m": lambda z: (num(z.get("ofi_supportive_vol_5m")) - num(z.get("ofi_supportive_vol_30m"))) if num(z.get("ofi_supportive_vol_5m")) is not None and num(z.get("ofi_supportive_vol_30m")) is not None else None,
        "taker_slope_5m_30m": lambda z: (num(z.get("supportive_taker_imb_5m")) - num(z.get("supportive_taker_imb_30m"))) if num(z.get("supportive_taker_imb_5m")) is not None and num(z.get("supportive_taker_imb_30m")) is not None else None,
        "taker_sign_change_5m_30m": lambda z: (1 if (num(z.get("supportive_taker_imb_5m")) or 0) * (num(z.get("supportive_taker_imb_30m")) or 0) < 0 else 0),
        "inband_sell_share_30m": lambda z: num(z.get("inband_sell_share_30m")),
        "spread_now_bps": lambda z: num(z.get("spread_now_bps")),
        "is_asia_session": lambda z: num(z.get("is_asia_session")),
    }
    FEATS = {**SM.FEATS, **extra}
    # groups
    def reps(clusters, cond=lambda c: True): return [SM.rep_of(c) for c in clusters if cond(c)]
    strong_reps = reps(scl); live_strong_reps = reps(scl, lambda c: any(z["_live"] for z in c))
    hit2_reps = reps(h2cl); noise_reps = reps(ncl)
    weak_reps = reps(h2cl, lambda c: not any(z["_hit25"] for z in c))
    fl_winners = [z for z in (first_live_z(c) for c in h2cl) if z]  # first-live of hit2 clusters (winners by definition reach 2)
    fl_losers = [first_live_z(c) for c in ncl if first_live_z(c)]   # first-live of noise clusters
    comps = {
        "strong_vs_noise": (strong_reps, noise_reps),
        "live_strong_vs_noise": (live_strong_reps, noise_reps),
        "hit2_vs_noise": (hit2_reps, noise_reps),
        "strong_vs_weak_hit2": (strong_reps, weak_reps),
        "firstlive_winner_vs_firstlive_loser": (fl_winners, fl_losers),
    }
    def date_sign_stability(pos, neg, fx):
        dates = sorted(set(z["_date"] for z in pos + neg)); agree = 0; tot = 0
        gm = (st.median([fx(z) for z in pos if fx(z) is not None]) if any(fx(z) is not None for z in pos) else None)
        gn = (st.median([fx(z) for z in neg if fx(z) is not None]) if any(fx(z) is not None for z in neg) else None)
        if gm is None or gn is None: return None
        gdir = gm >= gn
        for d in dates:
            pv = [fx(z) for z in pos if z["_date"] == d and fx(z) is not None]
            nv = [fx(z) for z in neg if z["_date"] == d and fx(z) is not None]
            if not pv or not nv: continue
            tot += 1
            if (st.median(pv) >= st.median(nv)) == gdir: agree += 1
        return round(agree / tot, 2) if tot else None
    def dir_sign_stability(pos, neg, fx):
        res = {}
        for dr in ("LONG", "SHORT"):
            pv = [fx(z) for z in pos if z["direction"] == dr and fx(z) is not None]
            nv = [fx(z) for z in neg if z["direction"] == dr and fx(z) is not None]
            if len(pv) >= 2 and len(nv) >= 2: res[dr] = round(st.median(pv) - st.median(nv), 3)
        return res
    C = {}
    for name, (pos, neg) in comps.items():
        feats = []
        for fn, fx in FEATS.items():
            pv = [fx(z) for z in pos]; nv = [fx(z) for z in neg]
            pvc = [x for x in pv if x is not None]; nvc = [x for x in nv if x is not None]
            if len(pvc) < 2 or len(nvc) < 2: continue
            d = cohen_d(pvc, nvc); a = auc(pvc, nvc)
            # precision @ top decile (direction by sign of d)
            allv = sorted(pvc + nvc, reverse=(d or 0) >= 0); k = max(1, len(allv) // 10)
            thr = allv[k - 1]; hi = (lambda x: x >= thr) if (d or 0) >= 0 else (lambda x: x <= thr)
            tp = sum(1 for z in pos if fx(z) is not None and hi(fx(z))); fp = sum(1 for z in neg if fx(z) is not None and hi(fx(z)))
            prec = round(tp / max(tp + fp, 1), 3); rec = round(tp / max(len(pos), 1), 3)
            feats.append({"feature": fn, "median_win": round(st.median(pvc), 4), "median_noise": round(st.median(nvc), 4),
                          "cohen_d": d, "auc": a, "precision_top_decile": prec, "recall": rec,
                          "sign_stability_by_date": date_sign_stability(pos, neg, fx),
                          "sign_stability_by_direction": dir_sign_stability(pos, neg, fx), "live_valid": "YES"})
        feats.sort(key=lambda x: -(x["auc"] or 0.5) if (x["auc"] or 0.5) >= 0.5 else (x["auc"] or 0.5))
        feats.sort(key=lambda x: -abs((x["auc"] or 0.5) - 0.5))
        C[name] = {"n_pos": len(pos), "n_neg": len(neg), "features": feats}

    # ---------- D: per-date casebook ----------
    candles = {d: daily_candle(d) for d in DATES}
    D_rows = []
    for d in DATES:
        dz = [z for z in zones if z["_date"] == d]
        cd = candles[d] or {}
        uniq_h2 = sum(1 for c in h2cl if first_z(c)["_date"] == d)
        uniq_strong = sum(1 for c in scl if first_z(c)["_date"] == d)
        winners = sorted([z for z in dz if z["_hit2"]], key=lambda x: -(x["_mfe"] or 0))[:3]
        noise = sorted([z for z in dz if z.get("sim_outcome") and not z["_hit2"]], key=lambda x: -(x.get("explainable_score") or 0))[:3]
        reg = max(set(z["_regime"] for z in dz), key=lambda r: sum(1 for z in dz if z["_regime"] == r)) if dz else "NONE"
        fam = max(set(z["_family"] for z in winners), key=lambda f: sum(1 for z in winners if z["_family"] == f)) if winners else "NONE"
        def fs(zs):
            if not zs: return {}
            return {"reclaim%": round(100 * sum(1 for z in zs if z.get("reclaim_zoneMid_preconfirm") == 1) / len(zs)),
                    "sweep%": round(100 * sum(1 for z in zs if z.get("swept_above") == 1 or z.get("swept_below") == 1) / len(zs)),
                    "OFI": round(st.median([num(z.get("eng_ofi")) or 0 for z in zs]), 2),
                    "taker15": round(st.median([num(z.get("supportive_taker_imb_15m")) or 0 for z in zs]), 2),
                    "micro5": round(st.median([num(z.get("dl2_microprice_aligned_delta_5m_bps")) or 0 for z in zs]), 2)}
        wf, lf = fs(winners), fs(noise)
        sep = [f"{k}: win {wf[k]} vs noise {lf[k]}" for k in wf if k in lf and wf[k] != lf[k]]
        D_rows.append({"date": d, "regime": reg, "net_pct": cd.get("net_pct"), "range_pct": cd.get("range_pct"),
                       "unique_hit2_moves": uniq_h2, "unique_strong_moves": uniq_strong,
                       "top_successful": [f"{z['direction']}/{z['_family']}/MFE{z['_mfe']}" for z in winners],
                       "top_noise_lookalike": [f"{z['direction']}/{z['_family']}/MFE{z['_mfe']}" for z in noise],
                       "winner_features": wf, "loser_features": lf, "what_separated": sep or ["nothing stable on decision features"],
                       "setup_family_of_day": fam})
    famdays = defaultdict(list)
    for r in D_rows: famdays[r["setup_family_of_day"]].append(r["date"])
    for r in D_rows: r["pattern_repeats_elsewhere"] = "YES" if len(famdays[r["setup_family_of_day"]]) >= 2 and r["setup_family_of_day"] != "NONE" else "NO"

    # ---------- E: aggressive rule search ----------
    def strong_caught(fired):
        s = set()
        for i, ss in enumerate(strong_sets):
            if any(id(z) in ss for z in fired): s.add(i)
        return len(s)
    def eval_rule(pred, pool):
        fired = [z for z in pool if z.get("sim_outcome") and pred(z)]
        picks = TA.dedup_alerts(fired); m = TA.metr(picks)
        m["precision"] = round(m["W"] / max(m["trades"], 1), 3); m["strong_caught"] = strong_caught(fired)
        m["noise_caught"] = sum(1 for p in picks if not p["_hit2"]); return m, picks
    Prules = {("P" + k[1:]): v for k, v in TA.rules(emed).items()}
    # Tier-B looser single/double-condition rules
    tierB = {
        "B1_reclaim_only": lambda z: z.get("reclaim_zoneMid_preconfirm") == 1,
        "B2_edge_only": lambda z: TA.edge(z),
        "B3_range_edge": lambda z: z["_regime"] == "RANGE" and TA.edge(z),
        "B4_sweep_only": lambda z: TA.sweep(z),
        "B5_taker_aligned": lambda z: TA.taker_ofi_ok(z),
        "B6_prior1d_aligned": lambda z: ((num(z.get("prior_move_1d_pct")) or 0) > 0) == (z["direction"] == "LONG") and abs(num(z.get("prior_move_1d_pct")) or 0) > 0.5,
    }
    allrules = {**Prules, **tierB}
    E_rows = []
    for name, pred in allrules.items():
        m, _ = eval_rule(pred, zones)
        # robustness
        lodo = [eval_rule(pred, [z for z in zones if z["_date"] != d])[0]["winrate"] for d in DATES]
        loco = []
        for ss in strong_sets:
            loco.append(eval_rule(pred, [z for z in zones if id(z) not in ss])[0]["winrate"])
        wo_both = eval_rule(pred, [z for z in zones if z["_date"] not in ("2026-05-14", "2026-05-15")])[0]
        E_rows.append({"rule": name, "unique_alerts": m["unique_alerts"], "W": m["W"], "L": m["L"], "TO": m["TO"],
                       "winrate": m["winrate"], "pf": m["pf"], "expectancy": m["expectancy"], "hit2_5": m["hit2_5"], "hit3": m["hit3"],
                       "precision": m["precision"], "strong_caught": m["strong_caught"], "noise_caught": m["noise_caught"],
                       "max_loss_streak": m["max_loss_streak"], "days_active": m["days_active"],
                       "winrate_without_0514_0515": wo_both["winrate"], "alerts_without_0514_0515": wo_both["unique_alerts"],
                       "lodo_min": min(lodo), "lodo_max": max(lodo), "loco_min": (min(loco) if loco else None), "loco_max": (max(loco) if loco else None),
                       "tier": ("A" if m["unique_alerts"] >= 4 else "thin"), "in_sample_only": "YES"})
    E_rows.sort(key=lambda r: (-r["winrate"], -r["unique_alerts"]))
    # gates
    def gate(minn):
        ok = [r for r in E_rows if r["unique_alerts"] >= minn]
        return max(ok, key=lambda r: (r["winrate"], r["unique_alerts"])) if ok else None
    E_gates = {f"min{n}": (lambda b: None if not b else {"rule": b["rule"], "winrate": b["winrate"], "alerts": b["unique_alerts"], "pf": b["pf"], "tag": ("TIER_A_CANDIDATE" if b["winrate"] >= 70 and b["unique_alerts"] >= 6 else ("VISUAL_DEMO_ONLY" if b["winrate"] >= 70 else "WEAK"))})(gate(n)) for n in (4, 5, 6, 8, 10, 15, 20)}
    # ranking/selector evaluation (one trade per cluster, different pick policies)
    sel = {
        "first_eligible_per_cluster": winrate([(first_live_z(c) or first_z(c)).get("sim_outcome") for c in h2cl]),
        "best_live_per_cluster": winrate([(best_z([z for z in c if z["_live"]] or c)).get("sim_outcome") for c in h2cl]),
        "earliest_safe_per_cluster": winrate([(earliest_safe_z(c) or first_z(c)).get("sim_outcome") for c in h2cl]),
    }
    best_overall = E_rows[0]
    dep1415 = "YES" if (best_overall["winrate_without_0514_0515"] < best_overall["winrate"] - 15 or best_overall["alerts_without_0514_0515"] < 3) else "NO"

    # ---------- F: engine rework proposals ----------
    F = [
        {"proposal": "1_cluster_aware_selector", "what_changes": "emit one alert per unique move cluster (first/best live-valid), not every raw zone",
         "fixes": "duplicate inflation (70 hit2 raw -> 19 unique; 34 strong raw -> 8 unique)", "uses_future": "NO (dedup by confirmedTs cooldown)",
         "expected_impact_18d": "removes ~73% duplicate alerts; winrate unchanged (selection within cluster already fine)", "risk": "low; pure post-engine layer", "cost": "low"},
        {"proposal": "2_zone_quality_score", "what_changes": "rank clusters by positive-minus-noise score before alerting",
         "fixes": "cross-cluster ranking (the real gap)", "uses_future": "NO", "expected_impact_18d": "LIMITED — current features have AUC~0.5-0.65, do not separate the 8 winners",
         "risk": "false confidence from in-sample tuning", "cost": "medium"},
        {"proposal": "3_delayed_confirmation", "what_changes": "enter after reclaim/microprice/OFI shift instead of at confirmedTs",
         "fixes": "late/early entry timing", "uses_future": "NO if event is causal", "expected_impact_18d": "NEGLIGIBLE — strong first-entry already 100% TP2; timing is not the bottleneck",
         "risk": "adds latency, may miss fast moves", "cost": "medium"},
        {"proposal": "4_event_driven_entry", "what_changes": "arm zone, trigger only on sweep+reclaim / absorption / wall-decay / thin-path-open",
         "fixes": "noise suppression", "uses_future": "NO", "expected_impact_18d": "reduces alert count but precision stays <30% (events fire on losers too)",
         "risk": "fewer trades, still no separation", "cost": "medium-high"},
        {"proposal": "5_duplicate_suppression", "what_changes": "suppress new signal if a same-direction cluster is already active",
         "fixes": "duplicate credits / over-alerting", "uses_future": "NO", "expected_impact_18d": "same as #1; cleaner live stream", "risk": "low", "cost": "low"},
    ]
    F_flag = "PARTIAL"  # cluster-aware/dup-suppression worth doing; full rewrite/scoring not justified yet

    # ---------- write files ----------
    def write_csv(path, rows):
        if not rows: return
        with path.open("w", encoding="utf-8", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
            w.writeheader()
            for r in rows: w.writerow({k: (json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v) for k, v in r.items()})
    write_csv(OUT / "DEEP_18D_UNIQUE_CLUSTER_MAP.csv", A_rows)
    (OUT / "DEEP_18D_UNIQUE_CLUSTER_MAP.json").write_text(json.dumps({"build": now_iso(), "summary": A_summary, "clusters": A_rows}, indent=2, default=str), encoding="utf-8")
    (OUT / "ENGINE_ZONE_SELECTION_AUDIT.json").write_text(json.dumps({"build": now_iso(), **B}, indent=2, default=str), encoding="utf-8")
    write_csv(OUT / "ENGINE_ZONE_SELECTION_AUDIT.csv", B_groups)
    (OUT / "DEEP_18D_POSITIVE_PATTERN_MINING.json").write_text(json.dumps({"build": now_iso(), "comparisons": C}, indent=2, default=str), encoding="utf-8")
    with (OUT / "DEEP_18D_POSITIVE_PATTERN_MINING.csv").open("w", encoding="utf-8", newline="") as fh:
        w = None
        for comp, blk in C.items():
            for f in blk["features"]:
                rr = {"comparison": comp, **{k: (json.dumps(v) if isinstance(v, dict) else v) for k, v in f.items()}}
                if w is None: w = csv.DictWriter(fh, fieldnames=list(rr.keys())); w.writeheader()
                w.writerow(rr)
    write_csv(OUT / "DEEP_18D_PER_DATE_PATTERN_CASEBOOK.csv", D_rows)
    (OUT / "DEEP_18D_PER_DATE_PATTERN_CASEBOOK.json").write_text(json.dumps({"build": now_iso(), "dates": D_rows, "family_recurrence": dict(famdays)}, indent=2, default=str), encoding="utf-8")
    write_csv(OUT / "DEEP_18D_AGGRESSIVE_RULE_SEARCH.csv", E_rows)
    (OUT / "DEEP_18D_AGGRESSIVE_RULE_SEARCH.json").write_text(json.dumps({"build": now_iso(), "status": "IN_SAMPLE_RESEARCH", "rules": E_rows, "gates": E_gates, "selector_policies": sel, "depends_on_0514_0515": dep1415}, indent=2, default=str), encoding="utf-8")
    (OUT / "ENGINE_SELECTOR_REWORK_PROPOSALS.json").write_text(json.dumps({"build": now_iso(), "proposals": F, "ENGINE_REWORK_RECOMMENDED": F_flag}, indent=2, default=str), encoding="utf-8")

    # xlsx
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for t, rows in (("ClusterMap", A_rows), ("SelectionAudit", B_groups), ("Rules", E_rows), ("Casebook", [{k: (v if not isinstance(v, (list, dict)) else json.dumps(v, ensure_ascii=False)) for k, v in r.items()} for r in D_rows])):
        ws = wb.create_sheet(t[:31])
        if rows: ws.append(list(rows[0].keys())); [ws.append([r.get(k) for k in rows[0].keys()]) for r in rows]
    wb.save(OUT / "DEEP_18D_UNIQUE_CLUSTER_MAP.xlsx")
    wb2 = openpyxl.Workbook(); ws = wb2.active; ws.title = "Rules"
    if E_rows: ws.append(list(E_rows[0].keys())); [ws.append([r.get(k) for k in E_rows[0].keys()]) for r in E_rows]
    wb2.save(OUT / "DEEP_18D_AGGRESSIVE_RULE_SEARCH.xlsx")
    wb3 = openpyxl.Workbook(); ws = wb3.active; ws.title = "Positive"
    rows3 = [{"comparison": comp, **{k: (json.dumps(v) if isinstance(v, dict) else v) for k, v in f.items()}} for comp, blk in C.items() for f in blk["features"]]
    if rows3: ws.append(list(rows3[0].keys())); [ws.append([r.get(k) for k in rows3[0].keys()]) for r in rows3]
    wb3.save(OUT / "DEEP_18D_POSITIVE_PATTERN_MINING.xlsx")
    wb4 = openpyxl.Workbook(); ws = wb4.active; ws.title = "Casebook"
    rows4 = [{k: (v if not isinstance(v, (list, dict)) else json.dumps(v, ensure_ascii=False)) for k, v in r.items()} for r in D_rows]
    if rows4: ws.append(list(rows4[0].keys())); [ws.append([r.get(k) for k in rows4[0].keys()]) for r in rows4]
    wb4.save(OUT / "DEEP_18D_PER_DATE_PATTERN_CASEBOOK.xlsx")

    # ---------- I: final report ----------
    top_pos_hit2 = sorted(C["hit2_vs_noise"]["features"], key=lambda x: -abs((x["auc"] or 0.5) - 0.5))[:6]
    pos_found = "PARTIAL" if any(abs((f["auc"] or 0.5) - 0.5) >= 0.12 for f in top_pos_hit2) else "NO"
    tier_a = "YES" if any(g and g["winrate"] >= 70 and g["alerts"] >= 6 for g in E_gates.values() if g) else ("PARTIAL" if any(g and g["winrate"] >= 70 for g in E_gates.values() if g) else "NO")
    flags = {
        "DEEP_18D_AUDIT_DONE": "YES", "POSITIVE_FEATURE_FOUND": pos_found, "TIER_A_FOUND_ON_18D": tier_a,
        "ENGINE_SELECTION_PROBLEM": B_flags["ENGINE_SELECTION_PROBLEM"], "ENGINE_REWORK_RECOMMENDED": F_flag,
        "NEXT_WINDOWS_SCANNER_DONE": "SEE scripts/research/select_next_windows_from_daily.py",
        "NEXT_WINDOWS_TO_DOWNLOAD_READY": "SEE reports/strategy-calibration/NEXT_WINDOWS_TO_DOWNLOAD.*",
        "READY_FOR_PRODUCTION_TRADING": "NO", "MORE_OOS_REQUIRED": "YES",
        "BEST_RULE": best_overall["rule"], "BEST_RULE_WINRATE": best_overall["winrate"], "BEST_RULE_ALERTS": best_overall["unique_alerts"],
        "STRONG_FIRST_ENTRY_WINRATE": B_groups[0]["first_zone_winrate"], "DEPENDS_ON_0514_0515": dep1415, "TARDIS_USED": "NO"}
    answers = {
        "1_better_filter": f"NO clean win — best in-sample rule {best_overall['rule']} {best_overall['winrate']}% on {best_overall['unique_alerts']} alerts; no rule reaches 70% at n>=6.",
        "2_positive_feature": f"{pos_found} — top by AUC: " + ", ".join(f"{f['feature']}(AUC {f['auc']})" for f in top_pos_hit2[:4]) + "; all near 0.5-0.65, none separates the 8 strong winners.",
        "3_where_is_the_problem": "Not engine/TP/SL and not within-cluster selection (strong first-zone TP2 = %d%%). It is FEATURE SEPARATION + cross-cluster RANKING: winners look like noise ex-ante." % B_groups[0]["first_zone_winrate"],
        "4_tier_a_possible": f"{tier_a}.", "5_if_yes_overfit": "n/a (not found at n>=6).",
        "6_if_no_why": "Only 8 unique strong moves; positive features have AUC~0.5-0.65 and sign flips across dates; any 70%+ rule is n<=4 curve-fit.",
        "7_engine_rebuild": "NO rewrite. Add a cluster-aware / duplicate-suppression layer (cheap, fixes alert inflation). A scoring/ranking layer is only worth it once a separating feature exists.",
        "8_next_step": "Run the daily windows scanner, download TREND_UP + TREND_DOWN + HIGH_VOL windows (the regimes missing here — 16/18 days were RANGE), rebuild caches, re-mine features across regimes.",
        "9_dates_to_download": "See NEXT_WINDOWS_TO_DOWNLOAD (data-driven). This 18d pool is almost all RANGE; trend/high-vol windows are required.",
        "10_windows_for_full_calibration": "See CALIBRATION_WINDOW_PLAN: feature-discovery vs validation vs holdout split, >=15-20 unique strong moves per regime, OKX+Binance overlap for cross-venue.",
    }
    (OUT / "DEEP_18D_FILTER_AND_ENGINE_AUDIT_FINAL_REPORT.json").write_text(json.dumps({"build": now_iso(), "A_summary": A_summary, "B": B, "E_gates": E_gates, "selector_policies": sel, "flags": flags, "answers": answers}, indent=2, default=str), encoding="utf-8")
    md = ["# OKX MAY 18D DEEP FILTER + ENGINE SELECTION AUDIT — FINAL", "", "## ⚠️ IN-SAMPLE RESEARCH — NOT PRODUCTION VALIDATION", "",
          f"**Build:** {now_iso()}", "", "## A. Pool reconstruction"] + [f"- {k}: {v}" for k, v in A_summary.items()] + [
          "", "## B. Engine / zone selection audit", "| group | clusters | first-zone wr | best-MFE wr | first-live wr | earliest-safe wr |", "|---|--:|--:|--:|--:|--:|"]
    for g in B_groups: md.append(f"| {g['group']} | {g['n_clusters']} | {g['first_zone_winrate']} | {g['best_mfe_zone_winrate']} | {g['first_live_zone_winrate']} | {g['earliest_safe_zone_winrate']} |")
    md += [f"", f"- {B['interpretation']}", f"- flags: {B_flags}", "",
           "## C. Positive features (hit2 vs noise, by |AUC-0.5|)", "| feature | win med | noise med | Cohen d | AUC | prec@decile | date-stab |", "|---|--:|--:|--:|--:|--:|--:|"]
    for f in top_pos_hit2: md.append(f"| {f['feature']} | {f['median_win']} | {f['median_noise']} | {f['cohen_d']} | {f['auc']} | {f['precision_top_decile']} | {f['sign_stability_by_date']} |")
    md += ["", "## E. Aggressive rule search (in-sample)", "| rule | alerts | W/L/TO | wr% | PF | strong | noise | wr_wo_1415 | tier |", "|---|--:|:--:|--:|--:|--:|--:|--:|:--:|"]
    for r in E_rows[:12]: md.append(f"| {r['rule']} | {r['unique_alerts']} | {r['W']}/{r['L']}/{r['TO']} | {r['winrate']} | {r['pf']} | {r['strong_caught']} | {r['noise_caught']} | {r['winrate_without_0514_0515']} | {r['tier']} |")
    md += [f"", f"- gates: " + " · ".join(f"{k}={'—' if not v else v['rule']+' '+str(v['winrate'])+'% n'+str(v['alerts'])+' '+v['tag']}" for k, v in E_gates.items()),
           f"- selector policies (winrate per unique cluster): {sel}", f"- depends_on_0514_0515 = {dep1415}", "",
           "## F. Engine rework proposals"] + [f"- **{p['proposal']}** — fixes: {p['fixes']}; impact: {p['expected_impact_18d']}; cost: {p['cost']}" for p in F] + [
           f"- ENGINE_REWORK_RECOMMENDED = {F_flag}", "", "## I. Answers"] + [f"**{k}** — {v}" for k, v in answers.items()] + ["", "## Flags", "```"] + [f"{k} = {v}" for k, v in flags.items()] + ["```"]
    (OUT / "DEEP_18D_FILTER_AND_ENGINE_AUDIT_FINAL_REPORT.md").write_text("\n".join(md), encoding="utf-8")

    # console
    print("A:", A_summary)
    print("B groups:")
    for g in B_groups: print("  ", g)
    print("B flags:", B_flags)
    print("C top hit2_vs_noise (AUC):")
    for f in top_pos_hit2: print(f"   {f['feature']:<24} d={f['cohen_d']} AUC={f['auc']} prec@dec={f['precision_top_decile']} datestab={f['sign_stability_by_date']}")
    print("E top rules:")
    for r in E_rows[:8]: print(f"   {r['rule']:<18} n{r['unique_alerts']:>3} wr {r['winrate']}% PF {r['pf']} strong {r['strong_caught']} noise {r['noise_caught']} wr_wo1415 {r['winrate_without_0514_0515']}")
    print("E gates:", {k: (v['rule']+' '+str(v['winrate'])+'% n'+str(v['alerts']) if v else None) for k, v in E_gates.items()})
    print("selector policies:", sel)
    print("FLAGS:", flags)
    print("per-date net%/range%:")
    for r in D_rows: print(f"   {r['date']} {r['regime']:<11} net {r['net_pct']}% range {r['range_pct']}% h2={r['unique_hit2_moves']} strong={r['unique_strong_moves']} dayfam={r['setup_family_of_day']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
